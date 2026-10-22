"""엑셀 출력 모듈. output/result_YYYYMMDD.xlsx 에 6개 시트를 생성한다."""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from src.judge import (
    COMPLIANT,
    EXCLUDED_FROM_RATE,
    NON_COMPLIANT,
    VERDICT_BASELINE,
    VERDICT_EARLY,
    VERDICT_LATE,
    VERDICT_NO_RECORD,
    VERDICT_NOT_REPLACED,
    VERDICT_ONTIME,
    Judgement,
)
from src.load_master import MasterRow
from src.models import ParseError, UnmatchedFragment

_FILL_RED = PatternFill("solid", fgColor="FFC7CE")
_FILL_YELLOW = PatternFill("solid", fgColor="FFEB9C")
_FILL_GREEN = PatternFill("solid", fgColor="C6EFCE")
_FILL_GRAY = PatternFill("solid", fgColor="D9D9D9")
_FONT_GRAY = Font(color="808080")
_FONT_BOLD = Font(bold=True)

_VERDICT_FILL = {
    VERDICT_LATE: _FILL_RED,
    VERDICT_NOT_REPLACED: _FILL_RED,
    VERDICT_EARLY: _FILL_YELLOW,
    VERDICT_ONTIME: _FILL_GREEN,
    VERDICT_BASELINE: _FILL_GRAY,
}

GRADE_ORDER = ["A", "B", "C"]
VERDICT_ORDER = [
    VERDICT_EARLY,
    VERDICT_ONTIME,
    VERDICT_LATE,
    VERDICT_NOT_REPLACED,
    VERDICT_BASELINE,
    VERDICT_NO_RECORD,
]


def _date_str(d: dt.date | None) -> str:
    return d.isoformat() if d else ""


def _auto_width(ws: Worksheet, max_width: int = 60) -> None:
    for col_cells in ws.columns:
        length = 0
        letter = None
        for cell in col_cells:
            letter = getattr(cell, "column_letter", None)
            if letter is None:
                continue
            if cell.value is not None:
                length = max(length, min(len(str(cell.value)), max_width))
        if letter:
            ws.column_dimensions[letter].width = max(10, length + 2)


def export_result(
    out_path: Path,
    judgements: list[Judgement],
    master_rows: list[MasterRow],
    unmatched: list[UnmatchedFragment],
    parse_errors: list[ParseError],
    as_of: dt.date,
    coverage: dict,
) -> None:
    wb = openpyxl.Workbook()
    default_sheet = wb.active
    if default_sheet is not None:
        wb.remove(default_sheet)

    _sheet_summary(wb.create_sheet("요약"), judgements, master_rows, as_of, coverage)
    _sheet_detail(wb.create_sheet("상세"), judgements)
    _sheet_overdue(wb.create_sheet("미교체현황"), judgements)
    _sheet_unmatched(wb.create_sheet("미매칭"), unmatched)
    _sheet_parse_errors(wb.create_sheet("파싱오류"), parse_errors)
    _sheet_upcoming(wb.create_sheet("다음예정"), judgements)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def _sheet_summary(ws, judgements, master_rows, as_of, coverage) -> None:
    ws.append([f"분석 기준일(--date): {as_of.isoformat()}"])
    ws.append([])

    counts: dict[tuple[str, str], dict[str, int]] = {}
    for j in judgements:
        bucket = counts.setdefault((j.standard_line, j.grade), {v: 0 for v in VERDICT_ORDER})
        bucket[j.verdict] += 1

    header = ["호기", "등급"] + VERDICT_ORDER + ["준수율(%)", "1소스(확인필요) 비율(%)"]
    ws.append(header)
    for cell in ws[ws.max_row]:
        cell.font = _FONT_BOLD

    need_check_by_group: dict[tuple[str, str], list[int]] = {}
    for j in judgements:
        g = need_check_by_group.setdefault((j.standard_line, j.grade), [0, 0])
        if j.verdict not in EXCLUDED_FROM_RATE:
            g[1] += 1
            if j.need_check == "Y":
                g[0] += 1

    for (line, grade) in sorted(counts.keys()):
        b = counts[(line, grade)]
        compliant = sum(b[v] for v in COMPLIANT)
        noncompliant = sum(b[v] for v in NON_COMPLIANT)
        denom = compliant + noncompliant
        rate = round(compliant / denom * 100, 1) if denom else ""
        nc = need_check_by_group.get((line, grade), [0, 0])
        nc_rate = round(nc[0] / nc[1] * 100, 1) if nc[1] else ""
        ws.append([line, grade] + [b[v] for v in VERDICT_ORDER] + [rate, nc_rate])

    ws.append([])
    judgeable = sum(1 for j in judgements if j.verdict not in EXCLUDED_FROM_RATE)
    total_items = len({(m.standard_line, m.item) for m in master_rows})
    ws.append(["판정 가능 항목 수 / 전체 항목 수", f"{judgeable} / {total_items}"])

    ws.append([])
    ws.append(["데이터 커버리지"])
    ws[ws.max_row][0].font = _FONT_BOLD
    ws.append(["출처", "시작일", "종료일"])
    for src in ("생산팀", "전산"):
        info = coverage.get(src, {})
        ws.append([src, _date_str(info.get("start")), _date_str(info.get("end"))])
    if coverage.get("gap_warning"):
        ws.append([coverage["gap_warning"]])
        ws[ws.max_row][0].font = Font(color="C00000", bold=True)

    _auto_width(ws)


def _sheet_detail(ws, judgements) -> None:
    header = [
        "호기", "공정(설비)", "교체 항목", "등급", "주기(일)",
        "직전교체일", "예상교체일", "실제교체일", "편차_일",
        "판정", "출처", "확인필요", "원문",
    ]
    ws.append(header)
    for cell in ws[1]:
        cell.font = _FONT_BOLD

    for j in sorted(judgements, key=lambda x: (x.standard_line, x.process, x.master_item, x.actual_date or dt.date.min)):
        ws.append([
            j.standard_line, j.process, j.master_item, j.grade, j.cycle_days,
            _date_str(j.prev_date), _date_str(j.expected_date), _date_str(j.actual_date),
            j.deviation_days if j.deviation_days is not None else "",
            j.verdict, j.source, j.need_check, j.raw_text,
        ])
        row = ws[ws.max_row]
        fill = _VERDICT_FILL.get(j.verdict)
        if fill:
            for cell in row:
                cell.fill = fill
        if j.verdict == VERDICT_NO_RECORD:
            for cell in row:
                cell.font = _FONT_GRAY
    _auto_width(ws)


def _sheet_overdue(ws, judgements) -> None:
    header = [
        "호기", "공정(설비)", "교체 항목", "등급", "주기(일)",
        "직전교체일", "예상교체일", "초과일수",
    ]
    ws.append(header)
    for cell in ws[1]:
        cell.font = _FONT_BOLD
    overdue = [j for j in judgements if j.verdict == VERDICT_NOT_REPLACED]
    overdue.sort(key=lambda x: x.deviation_days or 0, reverse=True)
    for j in overdue:
        ws.append([
            j.standard_line, j.process, j.master_item, j.grade, j.cycle_days,
            _date_str(j.prev_date), _date_str(j.expected_date), j.deviation_days,
        ])
        for cell in ws[ws.max_row]:
            cell.fill = _FILL_RED
    _auto_width(ws)


def _sheet_unmatched(ws, unmatched) -> None:
    ws.append(["출처", "날짜", "호기", "원문"])
    for cell in ws[1]:
        cell.font = _FONT_BOLD
    for u in sorted(unmatched, key=lambda x: (x.source, x.date or dt.date.min, x.standard_line)):
        ws.append([u.source, _date_str(u.date), u.standard_line, u.raw_text])
    _auto_width(ws)


def _sheet_parse_errors(ws, parse_errors) -> None:
    ws.append(["출처", "위치", "사유", "원본값"])
    for cell in ws[1]:
        cell.font = _FONT_BOLD
    for e in parse_errors:
        ws.append([e.source, e.location, e.reason, e.raw])
    _auto_width(ws)


def _sheet_upcoming(ws, judgements) -> None:
    ws.append(["호기", "공정(설비)", "교체 항목", "등급", "주기(일)", "직전교체일", "다음 예상 교체일"])
    for cell in ws[1]:
        cell.font = _FONT_BOLD

    latest: dict[tuple[str, str], Judgement] = {}
    no_record: list[Judgement] = []
    for j in judgements:
        key = (j.standard_line, j.master_item)
        if j.verdict == VERDICT_NO_RECORD:
            no_record.append(j)
            continue
        if j.actual_date is None:
            continue
        cur = latest.get(key)
        if cur is None or (j.actual_date or dt.date.min) > (cur.actual_date or dt.date.min):
            latest[key] = j

    upcoming = []
    for j in latest.values():
        if j.actual_date is None:
            continue
        next_date = j.actual_date + dt.timedelta(days=j.cycle_days)
        upcoming.append((next_date, j))
    upcoming.sort(key=lambda t: t[0])
    for next_date, j in upcoming:
        ws.append([
            j.standard_line, j.process, j.master_item, j.grade, j.cycle_days,
            _date_str(j.actual_date), _date_str(next_date),
        ])

    if no_record:
        ws.append([])
        ws.append(["[기준일 미설정 — 기록없음]"])
        ws[ws.max_row][0].font = _FONT_BOLD
        seen = set()
        for j in no_record:
            key = (j.standard_line, j.master_item)
            if key in seen:
                continue
            seen.add(key)
            ws.append([j.standard_line, j.process, j.master_item, j.grade, j.cycle_days, "", ""])
            for cell in ws[ws.max_row]:
                cell.font = _FONT_GRAY
    _auto_width(ws)
