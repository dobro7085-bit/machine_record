"""03 마스터(소모품 교체주기 기준) 로더.

시트 '교체주기_기준' 만 사용한다. 키 = (공정(설비), 교체 항목).
호기 구분이 없는 6개 공정은 CP1/CP2/PP1/PP2 4개 호기로 전개한다.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import openpyxl


MASTER_SHEET = "교체주기_기준"

LINE_EXPANDED_PROCESSES = {
    "혼합기",
    "도장",
    "디바이더",
    "순간 건조기(F/D)",
    "극판 적재기",
    "극판 적재 로봇",
}

EXPANSION_LINES = ("CP1", "CP2", "PP1", "PP2")


@dataclass
class MasterRow:
    no: int
    grade: str
    process: str
    item: str
    cycle_days: int
    standard_line: str  # 전개된 호기 또는 SMES 단독 공정명


def _grade_of(raw_grade: str) -> str:
    g = str(raw_grade).strip().upper()
    return g[0] if g else ""


def load_master(path: Path) -> list[MasterRow]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if MASTER_SHEET not in wb.sheetnames:
        raise ValueError(f"마스터 파일에 '{MASTER_SHEET}' 시트가 없습니다: {path}")
    ws = wb[MASTER_SHEET]

    rows = list(ws.iter_rows(values_only=True))
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    col = {name: i for i, name in enumerate(header)}

    required = ["등급", "공정(설비)", "교체 항목", "교체 주기(일)"]
    for r in required:
        if r not in col:
            raise ValueError(f"마스터 헤더에 '{r}' 컬럼이 없습니다. 실제 헤더: {header}")

    result: list[MasterRow] = []
    for raw in rows[1:]:
        if raw is None or all(c is None for c in raw):
            continue
        process = str(raw[col["공정(설비)"]]).strip() if raw[col["공정(설비)"]] else ""
        item = str(raw[col["교체 항목"]]).strip() if raw[col["교체 항목"]] else ""
        if not process or not item:
            continue
        grade = _grade_of(str(raw[col["등급"]]) if raw[col["등급"]] is not None else "")
        cycle = raw[col["교체 주기(일)"]]
        cycle_days = int(float(cycle)) if isinstance(cycle, (int, float, str)) and str(cycle).strip() not in ("", "None") else 0
        no_val = raw[col["No"]] if "No" in col else None
        no = int(no_val) if isinstance(no_val, (int, float)) else 0

        lines = EXPANSION_LINES if process in LINE_EXPANDED_PROCESSES else (process,)
        for line in lines:
            result.append(
                MasterRow(
                    no=no,
                    grade=grade,
                    process=process,
                    item=item,
                    cycle_days=cycle_days,
                    standard_line=line,
                )
            )
    return result
