"""01 생산팀 교체이력 로더.

시트 '2601이후' 만 사용. 4개 호기 블록(CP1/CP2/PP1/PP2, 각 8컬럼)의
'부동 내용'·'기타' 자유텍스트 셀을 (생산일, 호기, 원문)으로 추출한다.
날짜 3형식(datetime/'YYMMDD'/정수)을 모두 date 로 파싱하고, '주' 행 날짜를 '야' 행으로 forward-fill 한다.
"""
from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from src.models import ParseError

PROD_SHEET = "2601이후"

HEADER_ROWS = 3
DATA_START_ROW = 4

BLOCK_START_COL = {"CP1": 3, "CP2": 11, "PP1": 19, "PP2": 27}

FREETEXT_OFFSETS = {"부동내용": 3, "기타": 7}

_YYMMDD_RE = re.compile(r"^\d{6}$")


@dataclass
class ProdCell:
    date: dt.date
    standard_line: str
    column: str
    raw_text: str


def parse_prod_date(value) -> dt.date | None:
    """생산팀 A열 날짜 3형식 파싱. 실패 시 None.

    - datetime -> date
    - 정수 260909 / 문자열 '260401' -> 20YY-MM-DD
    """
    if value is None:
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, (int, float)):
        s = str(int(value))
        return _from_yymmdd(s)
    if isinstance(value, str):
        s = value.strip()
        if _YYMMDD_RE.match(s):
            return _from_yymmdd(s)
    return None


def _from_yymmdd(s: str) -> dt.date | None:
    if not _YYMMDD_RE.match(s):
        return None
    year = 2000 + int(s[0:2])
    month = int(s[2:4])
    day = int(s[4:6])
    try:
        return dt.date(year, month, day)
    except ValueError:
        return None


def load_prod(path: Path) -> tuple[list[ProdCell], list[ParseError]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if PROD_SHEET not in wb.sheetnames:
        raise ValueError(f"생산팀 파일에 '{PROD_SHEET}' 시트가 없습니다: {path}")
    ws = wb[PROD_SHEET]

    cells: list[ProdCell] = []
    errors: list[ParseError] = []
    current_date: dt.date | None = None

    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row_idx < DATA_START_ROW:
            continue
        raw_date = row[0] if len(row) > 0 else None
        shift = row[1] if len(row) > 1 else None

        if raw_date is not None and str(raw_date).strip() != "":
            parsed = parse_prod_date(raw_date)
            if parsed is None:
                errors.append(
                    ParseError(
                        source="생산팀",
                        location=f"{PROD_SHEET}!A{row_idx}",
                        reason="날짜 파싱 실패",
                        raw=repr(raw_date),
                    )
                )
            else:
                current_date = parsed

        has_text = _row_has_freetext(row)
        if current_date is None:
            if has_text:
                errors.append(
                    ParseError(
                        source="생산팀",
                        location=f"{PROD_SHEET}!row{row_idx}",
                        reason="날짜 미확정 상태의 교체 텍스트",
                        raw=_first_freetext(row),
                    )
                )
            continue

        for line, start in BLOCK_START_COL.items():
            for offset in FREETEXT_OFFSETS.values():
                col_idx = start - 1 + offset
                if col_idx >= len(row):
                    continue
                value = row[col_idx]
                if isinstance(value, str) and value.strip():
                    cells.append(
                        ProdCell(
                            date=current_date,
                            standard_line=line,
                            column=get_column_letter(col_idx + 1),
                            raw_text=value.strip(),
                        )
                    )
        _ = shift  # 야 행도 동일 생산일로 귀속(forward-fill)되므로 별도 처리 불필요
    return cells, errors


def _row_has_freetext(row) -> bool:
    for start in BLOCK_START_COL.values():
        for offset in FREETEXT_OFFSETS.values():
            col_idx = start - 1 + offset
            if col_idx < len(row):
                v = row[col_idx]
                if isinstance(v, str) and v.strip():
                    return True
    return False


def _first_freetext(row) -> str:
    for start in BLOCK_START_COL.values():
        for offset in FREETEXT_OFFSETS.values():
            col_idx = start - 1 + offset
            if col_idx < len(row):
                v = row[col_idx]
                if isinstance(v, str) and v.strip():
                    return v.strip()
    return ""
